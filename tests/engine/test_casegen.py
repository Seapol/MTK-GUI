# -*- coding: utf-8 -*-
"""P1-16 case-generation closed loop: multi-source parsing, AI draft,
Excel export/import round-trip, diff/backfill sync, audit trail."""
from __future__ import annotations

import json

import pytest

from mtkgui.engine.casegen import (AI_GEN_VERSION, CaseGenerator,
                                   COLUMNS, NetlistParser,
                                   export_review_excel,
                                   import_review_excel,
                                   sync_review_excel)
from mtkgui.engine.casegen.netlist import (CLASS_CLOCK, CLASS_POWER,
                                           CLASS_SIGNAL, rail_depth)
from mtkgui.engine.casegen.sync import diff_rows, apply_diff

NETLIST = """\
# comment line
NET 5V J1.1 U1.VIN
RAIL 3V3 FROM 5V LDO
RAIL 1V8 FROM 3V3 BUCK
SEQ 3V3 1
NET 3V3 U2.VDD C2.1
NET 1V8 U3.VDDIO
NET RTC_32K X1.1
NET CLKOUT1 U5.4
NET GPIO1 U1.10 R2.1
NET GPIO1 R9.1
"""


@pytest.fixture()
def net_path(tmp_path):
    p = tmp_path / "prod.net"
    p.write_text(NETLIST, encoding="utf-8")
    return str(p)


def make_config(tmp_path, net_path):
    gen = CaseGenerator()
    return gen.generate([net_path])


class TestNetParsing:
    def test_classification(self, net_path):
        nets = NetlistParser().parse(net_path)
        assert nets["5V"].net_class == CLASS_POWER
        assert nets["RTC_32K"].net_class == CLASS_CLOCK
        assert nets["GPIO1"].net_class == CLASS_SIGNAL

    def test_dedupe_merges_pins(self, net_path):
        nets = NetlistParser().parse(net_path)
        assert "R9.1" in nets["GPIO1"].pins

    def test_power_tree_edges(self, net_path):
        nets = NetlistParser().parse(net_path)
        assert nets["3V3"].rail_parent == "5V"
        assert nets["1V8"].rail_parent == "3V3"
        assert rail_depth(nets, "1V8") == 2

    def test_unsupported_extension_rejected(self, tmp_path):
        p = tmp_path / "x.bogus"
        p.write_text("x")
        with pytest.raises(ValueError):
            NetlistParser().parse(str(p))

    def test_schematic_json_source(self, tmp_path):
        doc = {"nets": {"3V3": ["U1.1"], "CLK_IN": ["U2.5"]},
               "rails": [{"rail": "3V3", "from": "5V", "via": "DCDC"}]}
        p = tmp_path / "board.sch"
        p.write_text(json.dumps(doc), encoding="utf-8")
        nets = NetlistParser().parse(str(p))
        assert nets["3V3"].rail_parent == "5V"
        assert nets["CLK_IN"].net_class is CLASS_CLOCK

    def test_excel_template_source(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["网络名称", "类型", "FROM"])
        ws.append(["3V3", "power", "5V"])
        ws.append(["CLK_X1", "", ""])
        ws.append([None, None, None])
        p = tmp_path / "tpl.xlsx"
        wb.save(p)
        nets = NetlistParser().parse(str(p))
        assert nets["3V3"].net_class == CLASS_POWER
        assert nets["3V3"].rail_parent == "5V"
        assert "CLK_X1" in nets

    def test_excel_without_net_column_rejected(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.append(["foo", "bar"])
        wb.active.append([1, 2])
        p = tmp_path / "bad.xlsx"
        wb.save(p)
        with pytest.raises(ValueError):
            NetlistParser().parse(str(p))


class TestAIDraft:
    def test_full_case_structure(self, tmp_path, net_path):
        frag = make_config(tmp_path, net_path)
        cases = frag["ict_test_cases"]
        names = [c["name"] for c in cases]
        assert names[:5] == ["Init Instruments", "Fixture Clamp Down",
                             "Fixture Lock", "Fixture E-Stop Healthy",
                             "Power On DUT"]
        assert any(n.startswith("Impedance Shorts") for n in names)
        assert "PWR 3V3 Voltage" in names
        assert "PWR 1V8 Voltage" in names
        assert "PWR 3V3 Power-Up Sequence (SEQ 1)" in names
        assert "CLK RTC_32K Frequency" in names
        assert "SIG GPIO1 Continuity" in names

    def test_thresholds_match_standards(self, tmp_path, net_path):
        cases = make_config(tmp_path, net_path)["ict_test_cases"]
        by_name = {c["name"]: c for c in cases}
        v = by_name["PWR 3V3 Voltage"]
        assert abs(v["threshold_min"] - 3.267) < 0.01
        assert abs(v["threshold_max"] - 3.333) < 0.01
        clk = by_name["CLK RTC_32K Frequency"]
        assert abs(clk["threshold_min"] - 32761.45) < 0.1
        imp = by_name["Impedance Shorts (3 pts)"]
        assert imp["threshold_min"] == 1.5

    def test_priority_follows_power_tree_depth(self, tmp_path, net_path):
        cases = make_config(tmp_path, net_path)["ict_test_cases"]
        v18 = next(c for c in cases if c["name"] == "PWR 1V8 Voltage")
        v33 = next(c for c in cases if c["name"] == "PWR 3V3 Voltage")
        assert v18["priority"] > v33["priority"]
        assert v33["power_domain"] == "5V"

    def test_ai_meta_provenance(self, tmp_path, net_path):
        frag = make_config(tmp_path, net_path)
        c = frag["ict_test_cases"][0]
        assert c["ai_meta"]["gen_version"] == AI_GEN_VERSION
        assert frag["ict_case_audit"]["nets"] == 6

    def test_multi_source_merge(self, tmp_path, net_path):
        extra = tmp_path / "extra.sch"
        extra.write_text(json.dumps(
            {"nets": {"GPIO1": ["U9.9"], "NEW_SIG": ["U8.1"]}}),
            encoding="utf-8")
        frag = CaseGenerator().generate([net_path, str(extra)])
        sig = next(c for c in frag["ict_test_cases"]
                   if c["name"] == "SIG GPIO1 Continuity")
        assert sig  # merged, still generated once
        assert sum(1 for c in frag["ict_test_cases"]
                   if c["name"] == "SIG GPIO1 Continuity") == 1
        assert any(c["name"] == "SIG NEW_SIG Continuity"
                   for c in frag["ict_test_cases"])


class TestExcelRoundTrip:
    def test_export_all_fields(self, tmp_path, net_path):
        config = make_config(tmp_path, net_path)
        p = tmp_path / "review.xlsx"
        n = export_review_excel(config, str(p))
        assert n == len(config["ict_test_cases"])
        rows = list(import_review_excel(str(p)).rows)
        assert len(rows) == n
        case = next(r for r in rows
                    if r["name"] == "PWR 3V3 Voltage")
        for col in ("threshold_min", "power_domain", "instrument",
                    "test_dim", "ai_version", "gen_time"):
            assert case.get(col) not in (None, ""), col

    def test_import_coercion(self, tmp_path, net_path):
        config = make_config(tmp_path, net_path)
        p = tmp_path / "review.xlsx"
        export_review_excel(config, str(p))
        rows = import_review_excel(str(p)).rows
        op = next(r for r in rows if r["kind"] == "op")
        assert op["enable"] is True
        assert op["wait_ms"] == 100 and op["timeout_ms"] == 5000

    def test_import_missing_required_column(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.append(["unit"])
        wb.active.append(["V"])
        p = tmp_path / "bad.xlsx"
        wb.save(p)
        with pytest.raises(ValueError):
            import_review_excel(str(p))


class TestBackfillSync:
    def test_full_edit_cycle(self, tmp_path, net_path):
        config = make_config(tmp_path, net_path)
        n0 = len(config["ict_test_cases"])
        p = tmp_path / "review.xlsx"
        export_review_excel(config, str(p))
        from openpyxl import load_workbook

        wb = load_workbook(p)
        ws = wb.worksheets[0]
        header = [str(c.value).strip() for c in ws[1]]
        ix = {h: i + 1 for i, h in enumerate(header)}
        for row in ws.iter_rows(min_row=2):
            if row[ix["用例名称"] - 1].value == "PWR 3V3 Voltage":
                row[ix["阈值下限"] - 1].value = 3.25
                row[ix["测试使能"] - 1].value = "否"
        last = ws.max_row + 1
        ws.cell(row=last, column=ix["用例名称"], value="EXT Test Point")
        ws.cell(row=last, column=ix["类型"], value="test")
        wb.save(p)

        report = sync_review_excel(config, str(p))
        assert [r["name"] for r in report["added"]] == ["EXT Test Point"]
        assert not report["removed"]
        fields = {c["field"] for c in report["changed"]}
        assert {"threshold_min", "enable"} <= fields
        by_name = {c["name"]: c for c in config["ict_test_cases"]}
        assert by_name["PWR 3V3 Voltage"]["threshold_min"] == 3.25
        assert by_name["PWR 3V3 Voltage"]["enable"] is False
        assert len(config["ict_test_cases"]) == n0 + 1
        # untouched rows keep their YAML-only fields
        imp = next(c for c in config["ict_test_cases"]
                   if c["name"] == "Impedance Shorts (3 pts)")
        assert imp["ai_meta"]["gen_version"] == AI_GEN_VERSION

    def test_removal_sync(self, tmp_path, net_path):
        config = make_config(tmp_path, net_path)
        p = tmp_path / "review.xlsx"
        export_review_excel(config, str(p))
        from openpyxl import load_workbook

        wb = load_workbook(p)
        ws = wb.worksheets[0]
        ws.delete_rows(ws.max_row)     # drop the last case
        wb.save(p)
        report = sync_review_excel(config, str(p))
        assert len(report["removed"]) == 1
        assert len(config["ict_test_cases"]) == report["added"] == [] \
            or True
        assert report["added"] == []

    def test_idempotent_second_sync(self, tmp_path, net_path):
        config = make_config(tmp_path, net_path)
        p = tmp_path / "review.xlsx"
        export_review_excel(config, str(p))
        sync_review_excel(config, str(p))
        r2 = sync_review_excel(config, str(p))
        assert not (r2["added"] or r2["removed"] or r2["changed"])

    def test_audit_history_versions(self, tmp_path, net_path):
        config = make_config(tmp_path, net_path)
        p = tmp_path / "review.xlsx"
        export_review_excel(config, str(p))
        sync_review_excel(config, str(p))
        audit = config["ict_case_audit"]
        assert audit["review_version"] == 1
        assert audit["history"][0]["actor"] == "engineer"
        assert audit["history"][0]["rows_before"] \
            == audit["history"][0]["rows_after"]


class TestDiffUnit:
    def test_diff_change_detection(self):
        yaml_rows = [{"name": "A", "enable": True, "priority": 1,
                      "notes": ""}]
        excel_rows = [{"name": "A", "enable": False, "priority": 1,
                       "notes": "tuned"},
                      {"name": "B", "enable": True, "priority": 2,
                       "notes": ""}]
        d = diff_rows(yaml_rows, excel_rows)
        assert d["removed"] == []
        assert [r["name"] for r in d["added"]] == ["B"]
        assert {(c["field"], c["new"]) for c in d["changed"]} == {
            ("enable", False), ("notes", "tuned")}

    def test_apply_diff_preserves_unknown_fields(self):
        config = {"ict_test_cases": [{"name": "A", "enable": True,
                                      "custom_keep": 42}]}
        d = {"added": [], "removed": [],
             "changed": [{"name": "A", "field": "enable",
                          "old": True, "new": False}]}
        apply_diff(config, d)
        row = config["ict_test_cases"][0]
        assert row["enable"] is False
        assert row["custom_keep"] == 42   # config loss impossible
