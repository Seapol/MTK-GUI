# -*- coding: utf-8 -*-
"""Unit tests for the two-way sync + validation layer and the Excel
round-trip / publishing."""

from __future__ import annotations

import pytest
import yaml

from mtkgui.gui.yamlbuild.model import YamlBuildModel
from mtkgui.gui.yamlbuild.sync import (
    sync_model_to_yaml,
    sync_preview_to_model,
    validate_yaml_text,
)
from mtkgui.gui.yamlbuild.excel_io import (
    export_to_excel,
    import_from_excel,
)
from mtkgui.gui.yamlbuild.publish import (
    archive_copy,
    build_plan_document,
    compare_plans,
    parse_plan_name,
    plan_filename,
    publish,
)


def make_model(enabled=True) -> YamlBuildModel:
    """Return a fully-enabled model with every required field filled
    (a valid, publishable configuration)."""
    model = YamlBuildModel()
    if enabled:
        model.enable_all()
    model.set_params("design_input", {
        "product_id": "IMXRT700", "part_number": "MTK12345",
        "sw_version": "1.2.3", "hw_version": "A",
        "batch": "MP (Production)",
    })
    model.set_params("parse_ict", {"netlist_file": "design.net",
                                   "ict_test_file": ""})
    model.set_params("instruments", {
        "psu_visa": "TCPIP0::192.168.1.20::inst0",
        "daq_visa": "GPIB0::9::INSTR", "dmm_visa": "",
        "self_test": "true",
    })
    model.set_params("rails", {
        "on_voltage_v": "5.0", "current_limit_a": "1.0",
        "on_delay_ms": "100", "off_delay_ms": "200",
        "off_protection": "true",
        "sequence": "VDD:0.0\nVDDCORE:0.2",
        "voltage_tolerance_pct": "0.1",
        "impedance_min_ohm": "1.5",
        "sample_rate_hz": "1000",
        "pre_trigger_s": "-0.5", "post_trigger_s": "6.0",
        "anomaly_policy": "stop",
    })
    model.set_params("clocks", {"clocks": "CLK1:32768:0.1",
                                "stabilize_ms": "100",
                                "drift_check": "true",
                                "multi_domain_check": "false"})
    model.set_params("gpios", {"groups": "LED1:PA0:out:none",
                               "level_threshold_v": "1.5",
                               "exception_check": "true"})
    model.set_params("fct_build", {"flow_steps": "step1\nstep2",
                                   "yield_threshold_pct": "98.0",
                                   "exception_branch": "stop",
                                   "case_link": ""})
    return model


# ---------------------------------------------------------------------------
# two-way sync + validation
# ---------------------------------------------------------------------------


def test_model_to_yaml_then_back_is_stable():
    """model -> YAML -> model keeps the effective dict identical (no
    drift, no loop artifacts)."""
    model = make_model()
    model.set_params("clocks", {"clocks": "CLK1:32768:0.1",
                                "stabilize_ms": "100",
                                "drift_check": "true",
                                "multi_domain_check": "false"})
    text = sync_model_to_yaml(model)
    other = make_model()
    result = sync_preview_to_model(other, text)
    assert result.ok, result.errors
    assert other.to_effective_dict() == model.to_effective_dict()


def test_syntax_error_detected_with_line():
    """A YAML syntax error is reported with its line (red marker)."""
    model = make_model()
    result = validate_yaml_text("yaml_build:\n  modules: [", model)
    assert not result.ok
    assert any("syntax" in msg for msg, _ in result.errors)


def test_invalid_edit_never_touches_model():
    """An invalid hand edit is rejected: the model keeps its state
    (dirty data never enters the configuration)."""
    model = make_model()

    def semantic(model_dict: dict) -> dict:
        """Content snapshot without the volatile save timestamp (the
        two to_dict() calls may straddle a second boundary under
        load)."""
        return {k: v for k, v in model_dict.items() if k != "saved_at"}

    before = semantic(model.to_dict())
    bad = yaml.safe_load(model.to_effective_yaml())
    bad["yaml_build"]["modules"]["rails"]["on_voltage_v"] = 99
    result = sync_preview_to_model(
        model, yaml.safe_dump(bad, sort_keys=False))
    assert not result.ok
    assert semantic(model.to_dict()) == before


def test_missing_section_rejected():
    """A YAML without the yaml_build section is rejected."""
    model = make_model()
    result = validate_yaml_text("foo: bar", model)
    assert not result.ok
    assert any("yaml_build" in msg for msg, _ in result.errors)


def test_unknown_module_rejected():
    """Unknown module keys are reported (no silent invention)."""
    model = make_model()
    result = validate_yaml_text(
        "yaml_build:\n  modules:\n    bogus: {stage_index: 0}", model)
    assert not result.ok
    assert any("unknown module" in msg for msg, _ in result.errors)


# ---------------------------------------------------------------------------
# Excel round-trip
# ---------------------------------------------------------------------------


def test_excel_round_trip(tmp_path):
    """Export -> import restores every parameter and enable state."""
    model = make_model()
    model.set_enabled("clocks", False)
    model.set_params("rails", {
        "sequence": "VDD:0.0\nVDDCORE:0.2",
        "voltage_tolerance_pct": "0.2",
        "impedance_min_ohm": "1.5",
        "sample_rate_hz": "2000",
        "pre_trigger_s": "-0.5",
        "post_trigger_s": "6.0",
        "anomaly_policy": "stop",
    })
    path = str(tmp_path / "cfg.xlsx")
    assert export_to_excel(model, path) > 10

    other = make_model(enabled=False)
    report = import_from_excel(other, path)
    assert report.applied, report.errors
    assert report.errors == []
    # the state dicts embed a save timestamp - compare content
    assert other.to_effective_dict() == model.to_effective_dict()
    for key in ("clocks", "rails", "design_input"):
        assert other.is_enabled(key) == model.is_enabled(key)
        assert other.get_params(key) == model.get_params(key)


def test_excel_import_rejects_invalid(tmp_path):
    """Invalid values abort the whole import with row errors (no
    partial overwrites)."""
    model = make_model()
    path = str(tmp_path / "cfg.xlsx")
    export_to_excel(model, path)

    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["YamlBuild"]
    for row in ws.iter_rows(min_row=2):
        if row[2].value == "on_voltage_v":
            row[3].value = "999"          # above maximum 60
    wb.save(path)

    other = make_model(enabled=False)
    report = import_from_excel(other, path)
    assert not report.applied
    assert any("above maximum" in e for e in report.errors)


def test_excel_import_missing_sheet(tmp_path):
    """A workbook without the expected sheet is rejected cleanly."""
    from openpyxl import Workbook
    path = str(tmp_path / "other.xlsx")
    wb = Workbook()
    wb.save(path)
    report = import_from_excel(make_model(), path)
    assert not report.applied
    assert any("missing worksheet" in e for e in report.errors)


# ---------------------------------------------------------------------------
# publishing: naming / content / archive / compare
# ---------------------------------------------------------------------------


def test_plan_filename_rules():
    """The fixed naming: Plan_[Core ID]_[Part#]_build_[kind]_v.<ver>."""
    model = make_model()
    name = plan_filename(model, "draft")
    assert name == "Plan_IMXRT700_MTK12345_build_draft_v.1.0.0.yaml"
    final = plan_filename(model, "final")
    assert final == "Plan_IMXRT700_MTK12345_build_final_v.1.0.0.yaml"
    parsed = parse_plan_name(name)
    assert parsed == {"core": "IMXRT700", "part": "MTK12345",
                      "kind": "draft", "version": "1.0.0"}
    assert parse_plan_name("random.yaml") is None


def test_plan_filename_requires_design_input():
    """Publishing without Core ID / Part # is rejected (fixed naming
    cannot be composed)."""
    model = YamlBuildModel()
    model.enable_all()
    with pytest.raises(ValueError):
        plan_filename(model, "draft")


def test_publish_draft_and_final(tmp_path):
    """Draft and Final files are written with correct header flags
    and the dynamic build version embedded."""
    model = make_model()
    draft = publish(model, "draft", tmp_path)
    final = publish(model, "final", tmp_path)
    assert draft.name.endswith("build_draft_v.1.0.0.yaml")
    assert final.name.endswith("build_final_v.1.0.0.yaml")
    doc = yaml.safe_load(final.read_text(encoding="utf-8"))
    assert doc["plan"]["kind"] == "final"
    assert doc["plan"]["locked"] is True
    assert doc["plan"]["core_id"] == "IMXRT700"
    assert doc["plan"]["built_with"]           # dynamic version suffix
    assert "modules" in doc["yaml_build"]


def test_archive_copy_and_compare(tmp_path):
    """Archiving writes a traceable copy per project (tenant-isolated
    archive, dynamic build suffix); compare produces a diff."""
    model = make_model()
    draft = publish(model, "draft", tmp_path)
    final = publish(model, "final", tmp_path)
    archived = archive_copy(draft, tmp_path, model.project_key())
    assert archived.exists()
    assert model.project_key() in str(archived)
    assert archived.parent.parent.name == "yamlbuild"
    diff = compare_plans(draft, final)
    assert "kind" in diff and "+draft" not in diff
    assert "-    \"kind\": draft" in diff or "kind: draft" in diff
    assert "final" in diff
