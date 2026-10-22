# -*- coding: utf-8 -*-
"""M5 PART 1: ICT Test Plan output files (xlsx + evidence md)."""
from __future__ import annotations

import json

import openpyxl
import pytest

from mtkgui.ict.netparse.capacity import DEFAULT_LIMITS
from mtkgui.ict.netparse.cli import main
from mtkgui.ict.netparse.export_testplan import (
    CLOCK_HEADER,
    CONTROL_IO_HEADER,
    GND_HEADER,
    OTHERS_HEADER,
    POWER_HEADER,
    SHEET_CLOCK,
    SHEET_CONTROL_IO,
    SHEET_GND,
    SHEET_NAMES,
    SHEET_OTHERS,
    SHEET_POWER,
    SHEET_SUMMARY,
    build_testplan_markdown,
    export_testplan_pair,
    project_name_from_path,
)
from mtkgui.ict.netparse.model import NetRecord, ParseNetsResult
from mtkgui.ict.netparse.priority import rank_result


def _result() -> ParseNetsResult:
    result = ParseNetsResult(
        source="NET-96317_A.net", total=9, format="net")
    result.power = [
        NetRecord(name="VDD_3V3", members=["TP1.1", "U1.5"],
                  net_type="power", provenance="rule:power",
                  confidence="high", input_rail=True),
        NetRecord(name="VDD_1V8", members=["U2.4"],
                  net_type="power", provenance="rule:power",
                  confidence="medium"),
    ]
    result.clock = [
        NetRecord(name="CLK_24M", members=["U1.10", "U4.1"],
                  net_type="clock_single", provenance="rule:clock",
                  confidence="high"),
    ]
    result.control_io = [
        NetRecord(name="GPIO_RST", members=["TP2.1"],
                  net_type="signal", provenance="control_io:rst",
                  confidence="high"),
    ]
    result.gnd = [
        NetRecord(name="GND", members=["J1.2", "U1.1"],
                  net_type="gnd_ref", sub_class="pcb",
                  provenance="gnd", confidence="high"),
    ]
    result.others = [
        NetRecord(name="USB_P", members=["U1.30", "J2.1"],
                  net_type="signal", sub_class="unclassified",
                  provenance="fallback:unclassified", confidence="low"),
        NetRecord(name="SDA", members=["U1.40", "U5.2"],
                  net_type="signal", sub_class="interface",
                  provenance="exclude:protocol", confidence="low"),
    ]
    rank_result(result)
    result.power[0].input_rail = True
    return result


def _export(tmp_path, **kwargs):
    path = str(tmp_path / "plan.xlsx")
    export_testplan_pair(_result(), DEFAULT_LIMITS, path, "96317",
                         **kwargs)
    return path


def test_sheet_names_and_headers(tmp_path):
    workbook = openpyxl.load_workbook(_export(tmp_path))
    assert tuple(workbook.sheetnames) == SHEET_NAMES
    assert tuple(c.value for c in workbook[SHEET_POWER][1]) == POWER_HEADER
    assert tuple(c.value for c in workbook[SHEET_CLOCK][1]) == CLOCK_HEADER
    assert tuple(c.value for c in
                 workbook[SHEET_CONTROL_IO][1]) == CONTROL_IO_HEADER
    assert tuple(c.value for c in workbook[SHEET_GND][1]) == GND_HEADER
    assert tuple(c.value for c in
                 workbook[SHEET_OTHERS][1]) == OTHERS_HEADER


def test_power_sheet_columns_and_input_rail_first(tmp_path):
    workbook = openpyxl.load_workbook(_export(tmp_path))
    rows = list(workbook[SHEET_POWER].iter_rows(min_row=2,
                                                values_only=True))
    assert [row[0] for row in rows] == ["VDD_3V3", "VDD_1V8"]
    first = rows[0]
    assert first[0] == "VDD_3V3"
    assert first[1].startswith("TP1.1")
    assert first[2] in (None, "")          # R-GND Min stays for humans
    assert float(first[3]) == 3.3
    assert first[6] in (None, "")          # channels blank on first export


def test_clock_nominal_in_khz(tmp_path):
    workbook = openpyxl.load_workbook(_export(tmp_path))
    rows = list(workbook[SHEET_CLOCK].iter_rows(min_row=2,
                                                values_only=True))
    assert rows[0][0] == "CLK_24M"
    assert float(rows[0][2]) == 24000.0


def test_gnd_rank_and_others_grouping(tmp_path):
    workbook = openpyxl.load_workbook(_export(tmp_path))
    gnd = list(workbook[SHEET_GND].iter_rows(min_row=2,
                                             values_only=True))
    assert gnd[0][1] == "GND"
    assert gnd[0][2] == "pcb"
    others = list(workbook[SHEET_OTHERS].iter_rows(min_row=2,
                                                   values_only=True))
    # fixed sub-class order: interface before unclassified
    assert [row[0] for row in others] == ["interface", "unclassified"]


def test_markdown_evidence_and_power_tree(tmp_path):
    result = _result()
    text = build_testplan_markdown(
        result, DEFAULT_LIMITS, "96317",
        overrides={"VDD_3V3": {"input_rail": True}})
    assert "### VDD_3V3" in text
    assert "- confidence: high" in text
    assert "- input_rail: true" in text
    assert "## Power tree flows" in text
    assert "VDD_3V3 (INPUT)" in text
    assert "## Others" in text


def test_export_pair_writes_evidence_file(tmp_path):
    path = _export(tmp_path)
    from pathlib import Path
    md = Path(path).with_name("plan_evidence.md")
    assert md.exists()
    assert "## Power tree flows" in md.read_text(encoding="utf-8")


def test_overrides_win_and_restore_input_rail(tmp_path):
    path = str(tmp_path / "plan.xlsx")
    overrides = {
        "VDD_3V3": {"nominal_v": 3.4, "test_point": "TP9.1",
                "r_v_channel": "CH1", "daq_ai_channel": "AI7"},
        "VDD_1V8": {"input_rail": True},
    }
    export_testplan_pair(_result(), DEFAULT_LIMITS, path, "96317",
                         overrides=overrides)
    workbook = openpyxl.load_workbook(path)
    rows = list(workbook[SHEET_POWER].iter_rows(min_row=2,
                                                values_only=True))
    by_name = {row[0]: row for row in rows}
    assert by_name["VDD_3V3"][1] == "TP9.1"
    assert float(by_name["VDD_3V3"][3]) == 3.4
    assert by_name["VDD_3V3"][6] == "CH1"
    assert by_name["VDD_3V3"][7] == "AI7"


def test_project_name_from_path():
    assert project_name_from_path("NET-96317_A.net") == "96317"
    assert project_name_from_path("/tmp/board.net") == "board"


def test_cli_export_plan(tmp_path, capsys):
    net = tmp_path / "NET-96317_A.net"
    net.write_text("*SIGNAL* 3V3\nU1.5 TP1.1\n*SIGNAL* GND\nJ1.2 U1.1\n",
                   encoding="utf-8")
    out = tmp_path / "ICT Test Plan for 96317.xlsx"
    code = main([str(net), "--export-plan", str(out)])
    assert code == 0
    assert out.exists()
    md = out.with_name(out.stem + "_evidence.md")
    assert md.exists()
    assert "test plan written" in capsys.readouterr().out
    assert tuple(openpyxl.load_workbook(str(out)).sheetnames) == SHEET_NAMES


# --------------------------------------------- M5-FIX1: library pass-through
#: inlet evidence: P1 power jack + SW1 rated power switch
INLET_LIBRARY = {
    "P1": {"part_number": "DC-005", "description": "PWR JACK"},
    "SW1": {"part_number": "SS-12D00", "description": "5A@125VAC"},
}


def _inlet_result() -> ParseNetsResult:
    result = ParseNetsResult(source="inlet.net", total=2)
    result.power = [
        NetRecord(name="VBAT_PWR", members=["U9.1"],
                  net_type="power", provenance="rule:power",
                  confidence="high"),
        NetRecord(name="12V_IN", members=["P1.1", "SW1.2"],
                  net_type="power", provenance="rule:power",
                  confidence="high"),
    ]
    rank_result(result)
    return result


def _input_rails_of(md_path) -> list[str]:
    from pathlib import Path
    text = Path(md_path).read_text(encoding="utf-8")
    body = text.split("## Input rails")[1].split("## ")[0]
    return [line.strip().lstrip("- ") for line in body.splitlines()
            if line.strip().startswith("- ") and "none" not in line]


def test_library_recovers_input_rail_and_power_order(tmp_path):
    path = tmp_path / "inlet.xlsx"
    export_testplan_pair(_inlet_result(), DEFAULT_LIMITS, str(path),
                         "inlet", library=INLET_LIBRARY)
    rows = list(openpyxl.load_workbook(str(path))[SHEET_POWER]
                .iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in rows] == ["12V_IN", "VBAT_PWR"]
    md = path.with_name("inlet_evidence.md")
    assert _input_rails_of(md) == ["12V_IN"]
    assert "- input_rail: true" in md.read_text(encoding="utf-8")


def test_without_library_keeps_working(tmp_path):
    path = tmp_path / "nolib.xlsx"
    export_testplan_pair(_inlet_result(), DEFAULT_LIMITS, str(path),
                         "nolib")
    rows = list(openpyxl.load_workbook(str(path))[SHEET_POWER]
                .iter_rows(min_row=2, values_only=True))
    # the name heuristic ranks VBAT (stage 0) first - degraded but safe
    assert [row[0] for row in rows] == ["VBAT_PWR", "12V_IN"]
    assert _input_rails_of(path.with_name("nolib_evidence.md")) == []


def test_cli_export_plan_passes_library(tmp_path, capsys):
    net = tmp_path / "NET-97122_A.net"
    net.write_text("*SIGNAL* VBAT_PWR\nU9.1\n"
                   "*SIGNAL* 12V_IN\nP1.1 SW1.2\n", encoding="utf-8")
    library = tmp_path / "library.json"
    library.write_text(
        '{"P1": {"part_number": "DC-005", "description": "PWR JACK"},'
        ' "SW1": {"part_number": "SS-12D00", "description": "5A@125VAC"}}',
        encoding="utf-8")
    out = tmp_path / "ICT Test Plan for 97122.xlsx"
    code = main([str(net), "--library", str(library),
                 "--export-plan", str(out)])
    assert code == 0
    capsys.readouterr()
    rows = list(openpyxl.load_workbook(str(out))[SHEET_POWER]
                .iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == "12V_IN"


# ------------------------- M5-FIX1 regression: five-project export parity
PARITY_PROJECTS = [
    ("29420", "NET-29420_H.net", "library_29420.json"),
    ("94923", "NET-94923_A.net", "library_94923.json"),
    ("96317", "NET-96317_A.net", "library_96317.json"),
    ("97122", "NET-97122_A.net", "library_97122.json"),
    ("96624", "NET-96624_A.net", "library_96624.json"),
]
#: fix2 CLI JSON baseline: input_rails set + rank1 power net
PARITY_EXPECT = {
    "29420": ["VCC_12V0", "VCC_VBUS"],
    "94923": [],
    "96317": ["DC_5V_IN", "5V_SDA_PSW", "5V_USB0_OTG"],
    "97122": ["12V_IN"],
    "96624": ["VDD_12V"],
}


@pytest.mark.parametrize("pid,net_name,lib_name", PARITY_PROJECTS)
def test_five_project_export_parity(tmp_path, pid, net_name, lib_name):
    from pathlib import Path

    from mtkgui.ict.netparse import parse_nets_file
    root = Path(__file__).resolve().parents[2]
    net = root / "projects" / pid / net_name
    lib = root / "reports" / lib_name
    if not net.exists() or not lib.exists():
        pytest.skip("project fixture not available")
    library = json.loads(lib.read_text(encoding="utf-8"))
    result = parse_nets_file(str(net), library=library)
    out = tmp_path / f"ICT Test Plan for {pid}.xlsx"
    export_testplan_pair(result, DEFAULT_LIMITS, str(out), pid,
                         library=library)
    rails = _input_rails_of(out.with_name(
        f"ICT Test Plan for {pid}_evidence.md"))
    assert rails == PARITY_EXPECT[pid]
    rows = list(openpyxl.load_workbook(str(out))[SHEET_POWER]
                .iter_rows(min_row=2, values_only=True))
    if PARITY_EXPECT[pid]:
        assert rows[0][0] == PARITY_EXPECT[pid][0]
