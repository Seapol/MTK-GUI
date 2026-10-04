# -*- coding: utf-8 -*-
"""P2-4 GUI YAML<->Excel import/export unit tests (headless offscreen)."""
from __future__ import annotations

import os

import pytest
import yaml
from openpyxl import Workbook, load_workbook

from mtkgui.engine.casegen.excel_io import COLUMNS, _HEADER_CN, \
    export_review_excel
from mtkgui.gui.case_io_page import CaseIOPage
from mtkgui.gui.case_store import CaseStore

CASES = [
    {"name": "PWR 3V3 Voltage", "kind": "test", "enable": True,
     "wait_ms": 100, "timeout_ms": 5000, "unit": "V",
     "threshold_min": 3.267, "threshold_max": 3.333,
     "test_dim": "voltage", "priority": 1, "instrument": "DAQM",
     "power_domain": "SOURCE", "upstream": "SOURCE",
     "downstream": "—", "notes": ""},
    {"name": "SIG GPIO1 Continuity", "kind": "test", "enable": True,
     "wait_ms": 100, "timeout_ms": 5000, "unit": "Ω",
     "threshold_min": 1.5, "threshold_max": "—", "test_dim": "impedance",
     "priority": 3, "instrument": "DAQM"},
]
BASE_CFG = {"project": {"revision": "1.0"},
            "ict_test_cases": [dict(c) for c in CASES],
            "ict_case_audit": {"ai_version": "P1-16.1", "history": []}}


def _workbook_row(case: dict, extra_header: str | None = None):
    row = [case.get(c, "") for c in COLUMNS]
    if extra_header:
        row.append("note")
    return row


@pytest.fixture()
def env(tmp_path):
    yaml_path = tmp_path / "project.yaml"
    yaml_path.write_text(yaml.safe_dump(BASE_CFG, sort_keys=False,
                                        allow_unicode=True),
                         encoding="utf-8")
    store = CaseStore(snapshot_dir=str(tmp_path / "snaps"))
    page = CaseIOPage(str(yaml_path), store=store)
    page.interactive = False
    xlsx = str(tmp_path / "review.xlsx")
    export_review_excel(BASE_CFG, xlsx)
    return page, store, str(yaml_path), xlsx, tmp_path


# ------------------------------------------------------------- export
def test_export_full_17_columns(env):
    page, _store, _yaml, _xlsx, tmp_path = env
    out = str(tmp_path / "out.xlsx")
    page.export_edit.setText(out)
    assert page.on_export() == out
    ws = load_workbook(out).worksheets[0]
    assert [c.value for c in ws[1]] == \
        [_HEADER_CN[c] for c in COLUMNS], "17-column CN header"
    assert ws.max_row == len(CASES) + 1, "no row loss"


def test_export_auto_names_and_blocks_empty(env):
    page, store, yaml_path, _xlsx, tmp_path = env
    page.export_edit.setText("")
    out = page.on_export()
    assert out and "ICT_REVIEW_" in out and os.path.exists(out), "auto name"
    # empty case list -> blocked
    empty = tmp_path / "empty.yaml"
    empty.write_text(yaml.safe_dump({"ict_test_cases": []}),
                     encoding="utf-8")
    page2 = CaseIOPage(str(empty), store=store)
    page2.interactive = False
    page2.export_edit.setText(str(tmp_path / "x.xlsx"))
    assert page2.on_export() is None, "empty export intercepted"


# ------------------------------------------------------------- import
def test_import_applies_diff_and_updates_yaml(env):
    page, store, yaml_path, xlsx, _tmp = env
    wb = load_workbook(xlsx)
    ws = wb.worksheets[0]
    rev = {v: k for k, v in _HEADER_CN.items()}
    idx = {rev[ws.cell(row=1, column=i + 1).value]: i
           for i in range(len(COLUMNS))}
    ws.cell(row=2, column=idx["instrument"] + 1, value="DMM")
    ws.cell(row=2, column=idx["priority"] + 1, value=5)
    ws.cell(row=2, column=idx["threshold_max"] + 1, value=3.40)
    wb.save(xlsx)
    fired = []
    page.cases_applied.connect(fired.append)
    page.import_edit.setText(xlsx)
    report = page.on_import()
    assert report is not None and not report["added"] \
        and not report["removed"], "pure param diff"
    fields = {(c["name"], c["field"]) for c in report["changed"]}
    assert ("PWR 3V3 Voltage", "instrument") in fields
    assert ("PWR 3V3 Voltage", "priority") in fields
    assert ("PWR 3V3 Voltage", "threshold_max") in fields
    updated = store.load(yaml_path)
    case = updated["ict_test_cases"][0]
    assert case["instrument"] == "DMM" and case["priority"] == 5 \
        and case["threshold_max"] == 3.40, "precise YAML sync"
    assert store.history(updated)[-1]["source"] == "gui-import"
    assert fired, "cases_applied emitted"
    assert "gui-import" in page.log.toPlainText(), "traceability log"
    snaps = [f for f in __import__("os").listdir(store.snapshot_dir)
             if f.endswith(".yaml")]
    assert snaps, "auto snapshot after import"


def test_import_added_removed_and_dirty_column(env):
    page, store, yaml_path, xlsx, _tmp = env
    wb = load_workbook(xlsx)
    ws = wb.worksheets[0]
    ws.cell(row=1, column=len(COLUMNS) + 1, value="工程师批注")
    ws.cell(row=2, column=len(COLUMNS) + 1, value="keep")
    ws.cell(row=3, column=len(COLUMNS) + 1, value="keep")
    ws.cell(row=3, column=1, value="")  # blank name -> row dropped
    wb.save(xlsx)
    page.import_edit.setText(xlsx)
    report = page.on_import()
    assert report is not None
    assert "工程师批注" in report["ignored_columns"], "dirty col filtered"
    assert [r["name"] for r in report["removed"]] == \
        ["SIG GPIO1 Continuity"], "blank-name row = removed"
    updated = store.load(yaml_path)
    assert len(updated["ict_test_cases"]) == 1


def test_import_added_row(env):
    page, store, yaml_path, xlsx, _tmp = env
    wb = load_workbook(xlsx)
    ws = wb.worksheets[0]
    ws.append(["SIG NEW Continuity", "test", True, 100, 5000, "Ω",
               1.5, "—", 3, "—", "—", "—", "DAQM", "impedance",
               "新增", "P1-16.1", "2026-01-01T00:00:00Z"])
    wb.save(xlsx)
    page.import_edit.setText(xlsx)
    report = page.on_import()
    assert [r["name"] for r in report["added"]] == ["SIG NEW Continuity"]
    names = {c["name"] for c in store.load(yaml_path)["ict_test_cases"]}
    assert "SIG NEW Continuity" in names


def test_import_missing_file_and_empty_path(env):
    page, _store, _yaml, _xlsx, tmp_path = env
    page.import_edit.setText("")
    assert page.on_import() is None
    page.import_edit.setText(str(tmp_path / "ghost.xlsx"))
    assert page.on_import() is None


def test_import_dirty_workbook_intercepted(env):
    page, store, yaml_path, _xlsx, tmp_path = env
    bad = str(tmp_path / "bad.xlsx")
    wb = Workbook()
    wb.active.append(["用例名称"])  # missing required 类型 column
    wb.active.append(["X"])
    wb.save(bad)
    before = store.load(yaml_path)
    page.import_edit.setText(bad)
    assert page.on_import() is None, "dirty workbook blocked"
    assert store.load(yaml_path) == before, "YAML untouched on failure"


def test_import_locked_case_blocks_whole_import(env):
    page, store, yaml_path, xlsx, _tmp = env
    wb = load_workbook(xlsx)
    ws = wb.worksheets[0]
    rev = {v: k for k, v in _HEADER_CN.items()}
    idx = {rev[ws.cell(row=1, column=i + 1).value]: i
           for i in range(len(COLUMNS))}
    ws.cell(row=2, column=idx["notes"] + 1, value="human edit")
    wb.save(xlsx)
    cfg = store.load(yaml_path)
    store.set_locked(cfg, "PWR 3V3 Voltage", True)
    store.save(cfg, yaml_path)
    before = store.load(yaml_path)
    page.import_edit.setText(xlsx)
    assert page.on_import() is None, "locked conflict -> whole import blocked"
    after = store.load(yaml_path)
    assert after["ict_test_cases"][0].get("notes", "") == "", "unchanged"


# ------------------------------------------------------------- shell
def test_shell_registers_case_io_route(qapp):
    from mtkgui.gui.shell import MainWindow
    win = MainWindow(baseline_version="test")
    win.mount_default_routes()
    assert "case_io" in win.route_keys
    win.navigate("case_io")
    assert "case_io" in win.cached_pages
