# -*- coding: utf-8 -*-
"""P2-4 YAML<->Excel GUI import/export demo (headless-safe, rc=0).

Full closed loop offscreen: AI draft -> one-click export (17 columns)
-> simulated engineer edit of the workbook -> one-click import with
diff report (added/removed/changed/instrument/priority) -> YAML
updated + snapshot -> dirty workbook intercepted -> locked-case import
blocked -> audit trail.  Exit 0 = all checkpoints passed.
"""
from __future__ import annotations

import os
import sys
import tempfile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from mtkgui.engine.casegen.excel_io import COLUMNS, _HEADER_CN  # noqa: E402
from mtkgui.engine.casegen.generator import CaseGenerator  # noqa: E402
from mtkgui.gui.case_io_page import CaseIOPage  # noqa: E402
from mtkgui.gui.case_store import CaseStore  # noqa: E402

NETLIST = """
NET 3V3 U1.VDD C2.1
NET 1V8 U1.CORE
RAIL 1V8 FROM 3V3 LDO
NET CLK_32K U2.1
NET GPIO1 U1.10 R2.1
"""


def main() -> int:
    app = QApplication.instance() or QApplication(sys.argv)
    tmp = tempfile.mkdtemp(prefix="p2_4_demo_")
    nl_path = os.path.join(tmp, "net.txt")
    with open(nl_path, "w", encoding="utf-8") as fh:
        fh.write(NETLIST)

    # AI draft via untouched P1-16 casegen -> project YAML
    cfg = dict(CaseGenerator().generate([nl_path]))
    yaml_path = os.path.join(tmp, "project.yaml")
    CaseStore.save(cfg, yaml_path)
    store = CaseStore(snapshot_dir=os.path.join(tmp, "snaps"))
    page = CaseIOPage(yaml_path, store=store)
    page.interactive = False
    fired = []
    page.cases_applied.connect(fired.append)

    # 1. one-click export: 17 columns, full fields
    page.export_edit.setText(os.path.join(tmp, "review.xlsx"))
    out = page.on_export()
    assert out and os.path.exists(out), "export path"
    wb = load_workbook(out)
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    assert header == [_HEADER_CN[c] for c in COLUMNS], "17-column header"
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    names = {r[0] for r in rows}
    assert any(n.startswith("PWR 3V3") for n in names), "full fields"
    assert len(rows) == len(cfg["ict_test_cases"]), "no row loss"

    # 2. engineer edits the workbook: param change + instrument +
    #    priority move + one added row + one removed row + dirty column
    names_l = [r[0] for r in rows]
    target = next(n for n in names_l if n.startswith("PWR 3V3 Voltage"))
    clk = next(n for n in names_l if n.startswith("CLK"))
    sig = next(n for n in names_l if n.startswith("SIG"))
    rev_cn = {v: k for k, v in _HEADER_CN.items()}
    idx = {rev_cn[h]: i for h, i in ((h, i) for i, h in enumerate(header))}
    data = [list(r) for r in rows]
    for row in data:
        if row[0] == target:
            row[idx["threshold_max"]] = 3.40          # parameter change
            row[idx["instrument"]] = "DMM"            # instrument re-assign
            row[idx["priority"]] = 7                  # priority move
        if row[0] == sig:
            row[0] = None                              # drop -> removed
    data.append(["SIG NEW_GPIO Continuity", "test", True, 100, 5000,
                 "Ω", 1.5, "—", 3, "—", "—", "—", "DAQM", "impedance",
                 "人工新增", "P1-16.1", "2026-10-04T00:00:00Z"])
    data[0].append("工程师批注")                       # unknown column
    for row in data[1:]:
        row.append("ok")
    wb2 = load_workbook(out)
    ws2 = wb2.worksheets[0]
    ws2.delete_rows(2, ws2.max_row)
    for row in data:
        ws2.append(row)
    ws2.cell(row=1, column=len(COLUMNS) + 1, value="工程师批注")
    wb2.save(out)

    # 3. one-click import: diff report + YAML update + snapshot
    page.import_edit.setText(out)
    report = page.on_import()
    assert report is not None, "import passes"
    assert len(report["added"]) == 1, "added row"
    assert len(report["removed"]) == 1, "removed row"
    changed = {(c["name"], c["field"]) for c in report["changed"]}
    assert (target, "threshold_max") in changed, "param change"
    assert (target, "instrument") in changed, "instrument re-assign"
    assert (target, "priority") in changed, "priority move"
    assert "工程师批注" in report["ignored_columns"], "dirty column filtered"
    updated = store.load(yaml_path)
    names_after = {c["name"] for c in updated["ict_test_cases"]}
    assert "SIG NEW_GPIO Continuity" in names_after, "YAML updated"
    assert sig not in names_after, "removed row gone"
    tgt = next(c for c in updated["ict_test_cases"]
               if c["name"] == target)
    assert tgt["instrument"] == "DMM" and tgt["priority"] == 7 \
        and tgt["threshold_max"] == 3.40, "precise sync"
    assert fired and "gui-import" in page.log.toPlainText(), "trace"
    snaps = os.listdir(store.snapshot_dir)
    assert snaps, "auto snapshot after import"
    hist = store.history(updated)
    assert hist[-1]["source"] == "gui-import", "history entry"

    # 4. dirty workbook interception (missing required column)
    bad = os.path.join(tmp, "bad.xlsx")
    from openpyxl import Workbook
    wbb = Workbook()
    wbb.active.append(["用例名称"])  # no 类型 column -> ValueError
    wbb.active.append(["X"])
    wbb.save(bad)
    before = CaseStore.load(yaml_path)
    page.import_edit.setText(bad)
    assert page.on_import() is None, "dirty workbook blocked"
    assert CaseStore.load(yaml_path) == before, "YAML untouched"

    # 5. locked case blocks the whole import (workbook differs from YAML)
    wb3 = load_workbook(out)
    ws3 = wb3.worksheets[0]
    for r in range(2, ws3.max_row + 1):
        if ws3.cell(row=r, column=1).value == target:
            ws3.cell(row=r, column=idx["threshold_max"] + 1, value=3.55)
    wb3.save(out)
    saved = CaseStore.load(yaml_path)
    store.set_locked(saved, target, True)
    CaseStore.save(saved, yaml_path)
    page.import_edit.setText(out)
    assert page.on_import() is None, "locked conflict blocks import"
    after = CaseStore.load(yaml_path)
    tgt2 = next(c for c in after["ict_test_cases"]
                if c["name"] == target)
    assert tgt2["threshold_max"] == 3.40, "locked case value unchanged"

    print("[P2-4 case io demo] YAML<->Excel GUI loop OK — 17-column "
          "export, edited import with diff report (added/removed/"
          "changed/instrument/priority), dirty-column filter, dirty-"
          "workbook interception, locked-case guard, auto snapshot + "
          "audit trail all validated")
    return 0


if __name__ == "__main__":
    sys.exit(main())
