# -*- coding: utf-8 -*-
"""P1-16 closed-loop demo - full forward + reverse pipeline, headless:

    netlist -> AI YAML draft -> review Excel -> (simulated human edit)
    -> Excel backfill sync -> audited production YAML

    python -m mtkgui.engine.casegen_demo                 # default tmp dir
    python -m mtkgui.engine.casegen_demo --dir <workdir>

Exit code: 0 when the loop closes cleanly, 1 on any assertion miss.
"""
from __future__ import annotations

import argparse
import tempfile
from pathlib import Path

from .casegen import (CaseGenerator, export_review_excel,
                      import_review_excel, sync_review_excel)

NETLIST = """\
# demo production netlist (FRDM-IMX93 style)
NET 5V J1.1 U1.VIN C1.1
RAIL 3V3 FROM 5V LDO
RAIL 1V8 FROM 3V3 BUCK
SEQ 3V3 1
SEQ 1V8 2
NET 3V3 U2.VDD C2.1 R5.2
NET 1V8 U3.VDDIO C3.1
NET RTC_32K X1.1 U4.RTC
NET CLKOUT1 U5.4
NET GPIO1 U1.10 R2.1
NET SDA1 U6.3 R7.1
"""


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="ICT case closed-loop demo (P1-16)")
    ap.add_argument("--dir", default=None,
                    help="work directory (default: fresh temp dir)")
    args = ap.parse_args(argv)
    work = Path(args.dir or tempfile.mkdtemp(prefix="casegen_demo_"))
    work.mkdir(parents=True, exist_ok=True)

    lines: list[str] = []
    log = lines.append

    # 1. import source + AI draft -----------------------------------
    src = work / "demo.netlist"
    src.write_text(NETLIST, encoding="utf-8")
    config: dict = {}
    gen = CaseGenerator(log_fn=log)
    config.update(gen.generate([str(src)]))
    n_cases = len(config["ict_test_cases"])
    assert n_cases > 10, "AI draft should produce a full ICT suite"

    # 2. export review Excel ----------------------------------------
    xlsx = work / "ict_review.xlsx"
    exported = export_review_excel(config, str(xlsx))
    assert exported == n_cases
    log(f"[CASEGEN] review excel exported: {xlsx.name} "
        f"({exported} row(s))")

    # 3. simulated human review edits -------------------------------
    from openpyxl import load_workbook

    wb = load_workbook(xlsx)
    ws = wb.worksheets[0]
    header = [str(c.value).strip() for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(header)}
    edited = 0
    for row in ws.iter_rows(min_row=2):
        name = row[idx["用例名称"] - 1].value
        # disable one signal continuity case (human decision)
        if isinstance(name, str) and name.startswith("SIG GPIO1"):
            row[idx["测试使能"] - 1].value = "否"
            edited += 1
        # retune one voltage window (human decision)
        if isinstance(name, str) and name == "PWR 3V3 Voltage":
            row[idx["阈值下限"] - 1].value = 3.25
            row[idx["阈值上限"] - 1].value = 3.35
            edited += 1
    # human adds a new test case row
    ws.append([c if c != "enable" else True for c in
               [None] * len(header)])
    last = ws.max_row
    for key, val in (("用例名称", "PWR VBUS Voltage"), ("类型", "test"),
                     ("测试使能", "是"), ("单位", "V"),
                     ("阈值下限", 4.95), ("阈值上限", 5.05),
                     ("测试优先级", 1), ("测试仪器分配", "DAQM"),
                     ("测试维度", "voltage"), ("备注说明", "human added")):
        ws.cell(row=last, column=idx[key], value=val)
    edited += 1
    # human removes nothing here; save
    wb.save(xlsx)
    log(f"[CASEGEN] simulated human review: {edited} edit(s) applied")

    # 4. backfill sync ----------------------------------------------
    report = sync_review_excel(config, str(xlsx))
    log(f"[CASEGEN] backfill: added={[r['name'] for r in report['added']]} "
        f"removed={[r['name'] for r in report['removed']]} "
        f"changed={len(report['changed'])} field edit(s)")
    assert any(r["name"] == "PWR VBUS Voltage"
               for r in report["added"]), "human-added row must land"

    # 5. verify YAML state ------------------------------------------
    by_name = {c["name"]: c for c in config["ict_test_cases"]}
    assert "PWR VBUS Voltage" in by_name
    assert by_name["SIG GPIO1 Continuity"]["enable"] is False
    assert by_name["PWR 3V3 Voltage"]["threshold_min"] == 3.25
    audit = config["ict_case_audit"]
    assert audit["history"] and audit["final"]["rows"] == n_cases + 1
    log(f"[CASEGEN] production YAML updated; audit version "
        f"v{audit['review_version']} "
        f"({audit['final']['rows']} case(s), final)")

    # 6. idempotent second sync (no-op) ------------------------------
    report2 = sync_review_excel(config, str(xlsx))
    assert not report2["added"] and not report2["removed"] \
        and not report2["changed"], "second sync must be a no-op"
    log("[CASEGEN] second sync is a clean no-op (no dirty residue)")

    print("\n".join(lines))
    print("CASEGEN DEMO: CLOSED LOOP OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
