#!/usr/bin/env python3
"""Build the 'MTK Common Platform' introduction deck (10-min management overview).

Run:  ./mtk_gui/bin/python scripts/make_intro_ppt.py
Output: MTK_Common_Platform_Intro.pptx (repo root)
"""
import os

import fitz  # PyMuPDF - used for image dimension probing
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY   = RGBColor(0x0E, 0x24, 0x3D)
ACCENT = RGBColor(0x00, 0xC2, 0xA8)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT  = RGBColor(0xC7, 0xD3, 0xE1)
INK    = RGBColor(0x1B, 0x2B, 0x3A)
MUTED  = RGBColor(0x5A, 0x6B, 0x7C)
BOX_BG = RGBColor(0xF2, 0xF6, 0xFA)
BOX_LN = RGBColor(0xD5, 0xDF, 0xE9)
CHIP_BG = RGBColor(0xE6, 0xF7, 0xF4)

BLUE   = RGBColor(0x15, 0x65, 0xC0)   # instruments (DAQ973A)
GREEN  = RGBColor(0x2E, 0x9E, 0x4F)   # fixture
AMBER  = RGBColor(0xE8, 0x9C, 0x00)   # flash probes
PURPLE = RGBColor(0x7C, 0x4D, 0xFF)   # DUT

SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def rect(slide, x, y, w, h, fill, line=None, line_w=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    sp.shadow.inherit = False
    return sp


def txt(slide, x, y, w, h, text, size, color, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font='Arial'):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color
    r.font.bold = bold; r.font.name = font
    return tb


def header(slide, title, kicker=None):
    rect(slide, Inches(0.55), Inches(0.42), Inches(0.12), Inches(0.62), ACCENT)
    txt(slide, Inches(0.85), Inches(0.36), Inches(11.9), Inches(0.75),
        title, 30, NAVY, bold=True)
    if kicker:
        txt(slide, Inches(0.87), Inches(1.02), Inches(11.9), Inches(0.4),
            kicker, 14, MUTED)
    rect(slide, Inches(0.55), Inches(1.42), Inches(12.23), Pt(1.4), BOX_LN)


def anno_card(slide, x, y, w, h, color, title, desc):
    """Annotation card: colored edge + bold title + one plain-English line."""
    rect(slide, x, y, w, h, BOX_BG, line=BOX_LN, line_w=Pt(0.75))
    rect(slide, x, y, Inches(0.07), h, color)
    txt(slide, x + Inches(0.18), y + Inches(0.05), w - Inches(0.3), Inches(0.28),
        title, 12, NAVY, bold=True)
    txt(slide, x + Inches(0.18), y + Inches(0.33), w - Inches(0.3), h - Inches(0.38),
        desc, 10.5, INK)


# ================================================================ Slide 1
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, Inches(0.9), Inches(2.05), Inches(0.16), Inches(2.0), ACCENT)
txt(s, Inches(1.35), Inches(2.0), Inches(11.5), Inches(1.0),
    'MTK Common Platform', 46, WHITE, bold=True)
txt(s, Inches(1.37), Inches(3.15), Inches(11.5), Inches(0.55),
    'Manufacturing Test Kit (MTK GUI)', 24, ACCENT, bold=True)
txt(s, Inches(1.37), Inches(3.8), Inches(11.6), Inches(0.55),
    'All-in-one: from design to platform, from net to system', 18, WHITE, bold=True)
txt(s, Inches(1.37), Inches(4.35), Inches(11.6), Inches(0.5),
    'One common platform for all NXP Commercial Boards', 16, LIGHT)
txt(s, Inches(1.37), Inches(4.9), Inches(11.6), Inches(0.5),
    'ICT  \u2192  Flash FAT  \u2192  FCT  \u2192  Flash OOBE', 15, LIGHT)
rect(s, Inches(1.37), Inches(6.85), Inches(10.6), Pt(1.2), ACCENT)
txt(s, Inches(1.37), Inches(6.95), Inches(11.5), Inches(0.4),
    'Developed & deployed by NXP HWS-CC Team \u2014 all TE engineers  \u00b7  October 2026',
    13, LIGHT)

# ================================================================ Slide 1b: Agenda
s = prs.slides.add_slide(BLANK)
header(s, 'Agenda', '10 minutes \u00b7 6 topics')

agenda = [
    ('What is MTK GUI?',
     'Platform overview, main interface & Yaml Build tour'),
    ('Why MTK?',
     'Manual testing vs. LabVIEW / TestStand + PXI vs. MTK'),
    ('Development & Deployment Roadmap',
     'What is done, what is in debug, what comes next'),
    ('Hardware BoM & Investment',
     '\u2248 $15\u201319k per station, zero software license fees'),
    ('MTK vs. Traditional Manual Testing',
     'Coverage, precision and new capabilities'),
    ('Summary & Ask',
     'What we need to move to Phase 3 and production rollout'),
]
ay, ah, agap = Inches(1.7), Inches(0.78), Inches(0.08)
for i, (t, d) in enumerate(agenda):
    y = ay + i * (ah + agap)
    rect(s, Inches(0.9), y, Inches(11.5), ah, BOX_BG, line=BOX_LN, line_w=Pt(0.75))
    rect(s, Inches(0.9), y, Inches(0.07), ah, ACCENT)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.12), y + Inches(0.19),
                              Inches(0.4), Inches(0.4))
    circ.fill.solid(); circ.fill.fore_color.rgb = NAVY
    circ.line.fill.background(); circ.shadow.inherit = False
    txt(s, Inches(1.12), y + Inches(0.235), Inches(0.4), Inches(0.32),
        str(i + 1), 13, WHITE, bold=True, align=PP_ALIGN.CENTER)
    txt(s, Inches(1.72), y + Inches(0.09), Inches(10.5), Inches(0.32),
        t, 15, NAVY, bold=True)
    txt(s, Inches(1.72), y + Inches(0.42), Inches(10.5), Inches(0.3),
        d, 11.5, MUTED)

# ================================================================ Slide 2
s = prs.slides.add_slide(BLANK)
header(s, 'What is MTK GUI?',
       'One PC + one rack = the complete test station. MTK software runs it all \u2014 on every NXP commercial board.')

# Left: rack block diagram
img_h = Inches(4.77)
img_w = Inches(4.77 * 2804 / 2252)          # keep native aspect (≈5.94")
s.shapes.add_picture('.ppt_assets/rack_diagram.png',
                     Inches(0.45), Inches(1.55), img_w, img_h)

# Right: 6 plain-English annotation cards
cx, cw = Inches(6.6), Inches(6.28)
cards = [
    (NAVY,   '1 \u00b7 Host PC \u2014 the brain',
     'One computer runs the MTK software, drives every instrument and produces the reports.'),
    (BLUE,   '2 \u00b7 Instruments \u2014 measure & power',
     'Digital multimeter, waveform capture and a programmable power supply for the board.'),
    (ACCENT, '3 \u00b7 Test peripherals',
     'Wi-Fi router, Bluetooth, display, camera \u2014 used to verify the product\u2019s real functions.'),
    (AMBER,  '4 \u00b7 Flash / debug probes',
     'Load the firmware into the board (factory-test image, then shipping image).'),
    (GREEN,  '5 \u00b7 Test fixture (the bed)',
     'Holds the board; 128 spring-loaded pins touch every test point automatically \u2014 no manual probing.'),
    (PURPLE, '6 \u00b7 DUT \u2014 the product',
     'The board under test, connected to the whole rack through the fixture.'),
]
cy, ch, cgap = Inches(1.55), Inches(0.72), Inches(0.09)
for color, t, d in cards:
    anno_card(s, cx, cy, cw, ch, color, t, d)
    cy = cy + ch + cgap

# Product traits chips row
traits = [
    'Desktop app \u2014 zero setup',
    'Lightweight & universal',
    'One YAML per project',
    'Friendly UI \u2014 quick to learn',
    'Supervisor + Operator roles',
]
tw, tgap = Inches(2.37), Inches(0.09)
for i, t in enumerate(traits):
    tx = Inches(0.55) + i * (tw + tgap)
    rect(s, tx, Inches(6.45), tw, Inches(0.42), CHIP_BG, line=ACCENT, line_w=Pt(1))
    txt(s, tx + Inches(0.06), Inches(6.5), tw - Inches(0.12), Inches(0.32),
        t, 10.5, NAVY, bold=True, align=PP_ALIGN.CENTER)

# Bottom takeaway band
rect(s, Inches(0.45), Inches(6.98), Inches(12.43), Inches(0.4), RGBColor(0xE6, 0xF7, 0xF4),
     line=ACCENT, line_w=Pt(1))
txt(s, Inches(0.65), Inches(7.02), Inches(12.0), Inches(0.32),
    'Takeaway: one operator, one click \u2014 MTK orchestrates the entire rack from a single PC.',
    13, NAVY, bold=True)

# ================================================================ Slide 2b
s = prs.slides.add_slide(BLANK)
header(s, 'MTK GUI Main Interface',
       'One window runs the whole station \u2014 shown here in the Ocean theme (virtual demo run).')

img_w2, img_h2 = Inches(8.6), Inches(5.12)
s.shapes.add_picture('.ppt_assets/main_window_ocean.png',
                     Inches(0.45), Inches(1.6), img_w2, img_h2)

def badge(slide, cx, cy, num):
    ov = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.14), Inches(cy - 0.14),
                                Inches(0.28), Inches(0.28))
    ov.fill.solid(); ov.fill.fore_color.rgb = NAVY
    ov.line.color.rgb = WHITE; ov.line.width = Pt(1.25)
    ov.shadow.inherit = False
    txt(slide, Inches(cx - 0.14), Inches(cy - 0.115), Inches(0.28), Inches(0.24),
        num, 11, WHITE, bold=True, align=PP_ALIGN.CENTER)

# Number badges are plain shapes + textboxes -> fully editable in PowerPoint.
for cx, cy, num in [(2.52, 2.40, '1'), (6.55, 1.87, '2'), (2.60, 3.28, '3'),
                    (1.35, 3.81, '4'), (5.05, 3.81, '5'), (1.12, 5.25, '6'),
                    (4.66, 6.43, '7')]:
    badge(s, cx, cy, num)

mods = [
    ('1 \u00b7 Product Info & Run Control',
     'Part #, SN, Run / Stop, long-run & interval'),
    ('2 \u00b7 Overall Flow & Result',
     'Live PASS / FAIL verdict per stage'),
    ('3 \u00b7 Workflow / Equipment / Yaml Build',
     'Test development & execution in one tool'),
    ('4 \u00b7 ICT / FCT Test Cases & Consoles',
     'Measured vs. limits; serial / SSH consoles'),
    ('5 \u00b7 Power Rails Waveform',
     '12-rail ms capture, auto-graded'),
    ('6 \u00b7 Event Log',
     'ms timestamps, colored verdicts'),
    ('7 \u00b7 Status Bar',
     'Version, role / mode, progress, LEDs'),
]
mx, mw = Inches(9.25), Inches(3.55)
my, mh, mgap = Inches(1.6), Inches(0.68), Inches(0.05)
for t, d in mods:
    anno_card(s, mx, my, mw, mh, ACCENT, t, d)
    my = my + mh + mgap

# ================================================================ Slide 2c
s = prs.slides.add_slide(BLANK)
header(s, 'Yaml Build \u2014 From Design Files to a Test Plan',
       'Five guided steps turn design inputs into a runnable ICT + FCT project \u2014 no manual data entry.')

img_w2, img_h2 = Inches(8.6), Inches(5.12)
s.shapes.add_picture('.ppt_assets/yamlbuild_ocean.png',
                     Inches(0.45), Inches(1.6), img_w2, img_h2)

# Number badges are plain shapes + textboxes -> fully editable in PowerPoint.
for cx, cy, num in [(0.75, 3.42, '1'), (1.55, 3.42, '2'), (2.35, 3.42, '3'),
                    (3.45, 3.42, '4'), (5.25, 3.75, '5')]:
    badge(s, cx, cy, num)

steps = [
    ('1 \u00b7 Design Input',
     'Point to netlist / SPF / PDF design files'),
    ('2 \u00b7 Parse Nets for ICT',
     'Power nets & test points extracted automatically'),
    ('3 \u00b7 Channel Allocation',
     'Nets auto-mapped to instrument channels'),
    ('4 / 5 \u00b7 ICT & FCT Sequences',
     'Test workflows generated, ready to run'),
    ('Live YAML Preview',
     'Two-way sync with the diagram; Excel import / export'),
]
mx, mw = Inches(9.25), Inches(3.55)
my, mh, mgap = Inches(1.6), Inches(0.92), Inches(0.12)
for t, d in steps:
    anno_card(s, mx, my, mw, mh, ACCENT, t, d)
    my = my + mh + mgap

# ================================================================ Slide 3
s = prs.slides.add_slide(BLANK)
header(s, 'Why MTK?',
       'Same lab instruments \u2014 three very different ways to run them.')

rows = [
    ('', 'Manual testing (today)', 'LabVIEW / TestStand + PXI', 'MTK Common Platform'),
    ('Test coverage',
     'Depends on the operator \u2014 steps get skipped under pressure',
     'High, but every new product needs specialist programming',
     'Full flow built in: ICT \u2192 flash \u2192 FCT \u2192 OOBE, identical on every unit'),
    ('Speed & efficiency',
     'Slow \u2014 one operator fully occupied per board',
     'Fast, but long (expensive) integration per product',
     'One click \u2014 fully automated: drives instruments, peripherals & fixture, talks to the DUT console; timing precise to the millisecond'),
    ('Quality assurance',
     'Human variance; mistakes hard to trace',
     'Repeatable, if configured strictly',
     'Same criteria for every unit; ATE fixture means the operator never hunts for test points or setup \u2014 recording & comparison are automatic'),
    ('Reporting',
     'Hand-written Excel \u2014 slow and error-prone',
     'Possible, extra setup per product',
     'Professional CSV / HTML / PDF reports; test data auto-uploaded to SharePoint cloud \u2014 zero extra steps, MES-ready'),
    ('Cost of ownership',
     'Cheap tools, expensive labor',
     'Software licenses + PXI chassis \u2014 high $ per station',
     'No license fees; off-the-shelf USB/LAN instruments'),
]
tbl = s.shapes.add_table(len(rows), 4, Inches(0.55), Inches(1.6),
                         Inches(12.23), Inches(5.2)).table
for c, w in zip(tbl.columns, (Inches(2.1), Inches(3.1), Inches(3.3), Inches(3.73))):
    c.width = w
MTK_FILL = RGBColor(0xE6, 0xF7, 0xF4)
for r, row in enumerate(rows):
    tbl.rows[r].height = Inches(0.72 if r else 0.5)
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.margin_left = Inches(0.1); cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = NAVY if c < 3 else ACCENT
        elif c == 3:
            cell.fill.fore_color.rgb = MTK_FILL
        else:
            cell.fill.fore_color.rgb = WHITE if r % 2 else BOX_BG
        p = cell.text_frame.paragraphs[0]
        run = p.add_run(); run.text = val
        f = run.font; f.name = 'Arial'
        if r == 0:
            f.size = Pt(13); f.bold = True
            f.color.rgb = NAVY if c == 3 else WHITE
        elif c == 0:
            f.size = Pt(12); f.bold = True; f.color.rgb = NAVY
        elif c == 3:
            f.size = Pt(11.5); f.bold = True; f.color.rgb = INK
        else:
            f.size = Pt(11); f.color.rgb = MUTED

rect(s, Inches(0.55), Inches(6.98), Inches(12.23), Inches(0.4),
     RGBColor(0xE6, 0xF7, 0xF4), line=ACCENT, line_w=Pt(1))
txt(s, Inches(0.75), Inches(7.02), Inches(11.8), Inches(0.32),
    'Takeaway: Professional automation at a fraction of the cost \u2014 built by NXP, for NXP boards.',
    13, NAVY, bold=True)

# ================================================================ Slide 4
s = prs.slides.add_slide(BLANK)
header(s, 'Development & Deployment Roadmap',
       'FCT proven on a real board \u00b7 ICT plans ready for 8 board types \u2014 real-mode validation is the next gate.')

phase_data = [
    (GREEN, 'Phase 1 \u00b7 DONE', [
        'Foundation infra, architecture, GUI and operating modes',
        'Test workflow and core logic developed & verified end to end',
    ]),
    (ACCENT, 'Phase 2 \u00b7 IN DEV & DEBUG', [
        'FCT full flow on real FRDM-IMX93: acceptance passed, no obvious bugs',
        'ICT test plans built from design files of 8 different board types \u2014 build OK, virtual-mode verified',
        'Real-mode ICT validation pending \u2014 instruments not yet on site',
    ]),
    (RGBColor(0x8A, 0x9B, 0xAC), 'Phase 3 \u00b7 NEXT', [
        'Run the full flow in real mode \u2014 all scenarios, end to end',
        'Adapt more board types onto one common platform',
        'Platform performance optimization \u2014 higher development & production test efficiency',
        'Connect to factory cloud server / MES',
    ]),
]
pw, pgap = Inches(3.87), Inches(0.31)
py, ph = Inches(1.6), Inches(3.3)
for i, (color, title, bullets) in enumerate(phase_data):
    px = Inches(0.55) + i * (pw + pgap)
    rect(s, px, py, pw, ph, BOX_BG, line=BOX_LN, line_w=Pt(1))
    rect(s, px, py, pw, Inches(0.46), color)
    txt(s, px + Inches(0.16), py + Inches(0.06), pw - Inches(0.32), Inches(0.34),
        title, 14.5, WHITE if i != 1 else NAVY, bold=True)
    tb = s.shapes.add_textbox(px + Inches(0.16), py + Inches(0.58),
                              pw - Inches(0.32), ph - Inches(0.7))
    tf = tb.text_frame; tf.word_wrap = True
    for j, b in enumerate(bullets):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run(); r.text = '\u2022  ' + b
        r.font.size = Pt(11.5); r.font.color.rgb = INK; r.font.name = 'Arial'
    if i < 2:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                px + pw + Inches(0.02), py + Inches(1.5),
                                Inches(0.27), Inches(0.3))
        ar.fill.solid(); ar.fill.fore_color.rgb = ACCENT
        ar.line.fill.background(); ar.shadow.inherit = False
    if i == 1:  # YOU ARE HERE marker
        mk = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                px + pw - Inches(1.62), py - Inches(0.3),
                                Inches(1.62), Inches(0.3))
        mk.fill.solid(); mk.fill.fore_color.rgb = NAVY
        mk.line.fill.background(); mk.shadow.inherit = False
        mk.adjustments[0] = 0.5
        txt(s, px + pw - Inches(1.62), py - Inches(0.27), Inches(1.62), Inches(0.26),
            'WE ARE HERE', 10.5, ACCENT, bold=True, align=PP_ALIGN.CENTER)

# Challenges band
rect(s, Inches(0.55), Inches(5.15), Inches(12.23), Inches(0.85),
     RGBColor(0xFD, 0xF3, 0xE3), line=RGBColor(0xE8, 0x9C, 0x00), line_w=Pt(1))
txt(s, Inches(0.75), Inches(5.22), Inches(2.6), Inches(0.3),
    'Current challenges', 13, RGBColor(0xB4, 0x6E, 0x00), bold=True)
txt(s, Inches(0.75), Inches(5.53), Inches(11.9), Inches(0.42),
    'Instrument rack not yet available on site \u2014 real-mode ICT validation is blocked (the Phase-3 gate)  \u00b7  per-product flash-probe integration  \u00b7  Windows packaging & code signing',
    12, INK)

rect(s, Inches(0.55), Inches(6.98), Inches(12.23), Inches(0.4),
     RGBColor(0xE6, 0xF7, 0xF4), line=ACCENT, line_w=Pt(1))
txt(s, Inches(0.75), Inches(7.02), Inches(11.8), Inches(0.32),
    'Takeaway: the core is done and proven \u2014 from here it is hardening, rollout and scale.',
    13, NAVY, bold=True)

# ================================================================ Slide 5
s = prs.slides.add_slide(BLANK)
header(s, 'Hardware BoM & Investment',
       'One complete test station \u2014 all off-the-shelf instruments, no PXI chassis required.')

def add_image_fit(slide, path, x, y, w, h):
    """Insert picture centered inside the (x, y, w, h) box, aspect preserved."""
    p = fitz.Pixmap(path)
    ar = p.width / p.height
    box_ar = w / h
    if ar <= box_ar:
        fh = h; fw = int(h * ar)
    else:
        fw = w; fh = int(w / ar)
    slide.shapes.add_picture(path, x + int((w - fw) / 2), y + int((h - fh) / 2), fw, fh)


# Top: 5 instrument photo cards (Keysight official list prices)
instruments = [
    ('.ppt_assets/DAQ973A.png', 'DAQ973A mainframe', 'DMM \u00b7 3 slots', '$3,569'),
    ('.ppt_assets/DAQM908A.png', 'DAQM908A \u00d7 2', '40-ch mux each', '$1,824'),
    ('.ppt_assets/DAQM907A.png', 'DAQM907A', 'Fixture IO \u00b7 AO', '$895'),
    ('.ppt_assets/U2355A.png', 'U2355A \u2020', 'Rail capture', '$2,100'),
    ('.ppt_assets/N5747A.png', 'N5747A PSU (optional)', 'Use DUT adapter when provided', '$3,902'),
]
iw, igap = Inches(2.37), Inches(0.09)
iy, ih = Inches(1.58), Inches(3.05)
for i, (img, model, role, price) in enumerate(instruments):
    ix = Inches(0.55) + i * (iw + igap)
    rect(s, ix, iy, iw, ih, BOX_BG, line=BOX_LN, line_w=Pt(1))
    if img:
        add_image_fit(s, img, ix + Inches(0.1), iy + Inches(0.1),
                      iw - Inches(0.2), Inches(2.0))
    else:
        rect(s, ix + Inches(0.1), iy + Inches(0.1), iw - Inches(0.2), Inches(2.0),
             RGBColor(0xE9, 0xEE, 0xF3))
        txt(s, ix + Inches(0.1), iy + Inches(0.75), iw - Inches(0.2), Inches(0.4),
            model.split(' \u2020')[0].split(' PSU')[0], 20, RGBColor(0x9F, 0xAF, 0xBE),
            bold=True, align=PP_ALIGN.CENTER)
        txt(s, ix + Inches(0.1), iy + Inches(1.2), iw - Inches(0.2), Inches(0.3),
            'product photo pending', 10, RGBColor(0x9F, 0xAF, 0xBE),
            align=PP_ALIGN.CENTER)
    if 'optional' in model.lower():
        badge = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   ix + iw - Inches(1.0), iy + Inches(0.14),
                                   Inches(0.86), Inches(0.24))
        badge.fill.solid(); badge.fill.fore_color.rgb = AMBER
        badge.line.fill.background(); badge.shadow.inherit = False
        badge.adjustments[0] = 0.5
        txt(s, ix + iw - Inches(1.0), iy + Inches(0.155), Inches(0.86), Inches(0.22),
            'OPTIONAL', 9, WHITE, bold=True, align=PP_ALIGN.CENTER)
    txt(s, ix + Inches(0.1), iy + Inches(2.16), iw - Inches(0.2), Inches(0.26),
        model, 11.5, NAVY, bold=True, align=PP_ALIGN.CENTER)
    txt(s, ix + Inches(0.1), iy + Inches(2.41), iw - Inches(0.2), Inches(0.22),
        role, 9.5, MUTED, align=PP_ALIGN.CENTER)
    txt(s, ix + Inches(0.1), iy + Inches(2.66), iw - Inches(0.2), Inches(0.28),
        price, 13, RGBColor(0x00, 0x8A, 0x76), bold=True, align=PP_ALIGN.CENTER)

# Bottom band: investment summary
rect(s, Inches(0.55), Inches(4.85), Inches(12.23), Inches(1.75),
     BOX_BG, line=BOX_LN, line_w=Pt(1))
txt(s, Inches(0.85), Inches(5.05), Inches(5.6), Inches(0.7),
    '\u2248 $15\u201319k / station', 34, NAVY, bold=True)
txt(s, Inches(0.85), Inches(5.85), Inches(5.6), Inches(0.3),
    'One-time hardware \u00b7 $0 software license (in-house)', 12, INK)
rect(s, Inches(6.55), Inches(5.05), Pt(1), Inches(1.35), BOX_LN)
txt(s, Inches(6.95), Inches(5.1), Inches(5.5), Inches(0.3),
    'vs. PXI alternative', 13, NAVY, bold=True)
txt(s, Inches(6.95), Inches(5.45), Inches(5.6), Inches(0.65),
    'Typical PXI station: $30\u201350k+ plus annual software licenses. MTK runs on USB/LAN instruments you may already own.',
    11.5, INK)
txt(s, Inches(0.85), Inches(6.28), Inches(11.6), Inches(0.22),
    '*Keysight US list prices (Oct 2026) \u00b7 \u2020U2355A discontinued \u2014 reseller price \u00b7 ~ = estimate',
    8.5, MUTED)

# ================================================================ Slide 5b
s = prs.slides.add_slide(BLANK)
header(s, 'Investment & Deployment Strategy',
       'Per-station total \u2248 $15.4\u201319.3k (PSU optional) \u2014 one BoM, two use cases.')

rest = [
    ('Other items', 'Qty', 'What it does', 'Price*'),
    ('J-Link flash probe', '1', 'Firmware flashing (FAT / OOBE) \u2014 reusable', '~$500'),
    ('HOST PC (Windows 11 Pro)', '1', 'Runs the MTK desktop app \u2014 reusable', '~$1,500'),
    ('ATE test fixture', '1', 'Custom per project \u2014 incl. assembly labor', '~$4,000+'),
    ('Peripherals & cabling', '1 set', 'Custom per project \u2014 incl. assembly labor', '~$1,000+'),
    ('TOTAL PER STATION', '', 'PSU optional \u2014 DUT adapter first',
     '\u2248 $15.4\u201319.3k'),
]
tbl = s.shapes.add_table(len(rest), 4, Inches(0.55), Inches(1.6),
                         Inches(12.23), Inches(2.7)).table
for c, w in zip(tbl.columns, (Inches(3.1), Inches(0.8), Inches(5.7), Inches(2.63))):
    c.width = w
for r, row in enumerate(rest):
    tbl.rows[r].height = Inches(0.45)
    last = (r == len(rest) - 1)
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.margin_left = Inches(0.1); cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = NAVY
        elif last:
            cell.fill.fore_color.rgb = ACCENT
        else:
            cell.fill.fore_color.rgb = WHITE if r % 2 else BOX_BG
        p = cell.text_frame.paragraphs[0]
        if c == 1:
            p.alignment = PP_ALIGN.CENTER
        if c == 3:
            p.alignment = PP_ALIGN.RIGHT
        run = p.add_run(); run.text = val
        f = run.font; f.name = 'Arial'
        if r == 0:
            f.size = Pt(12.5); f.bold = True; f.color.rgb = WHITE
        elif last:
            f.size = Pt(13); f.bold = True; f.color.rgb = NAVY
        elif c == 0:
            f.size = Pt(11.5); f.bold = True; f.color.rgb = INK
        elif c == 3:
            f.size = Pt(12); f.bold = True; f.color.rgb = INK
        else:
            f.size = Pt(11); f.color.rgb = MUTED

# Two deployment strategy cards
def deploy_card(x, color, title, lines):
    rect(s, x, Inches(4.55), Inches(6.0), Inches(2.0), BOX_BG, line=BOX_LN, line_w=Pt(1))
    rect(s, x, Inches(4.55), Inches(0.07), Inches(2.0), color)
    txt(s, x + Inches(0.24), Inches(4.71), Inches(5.5), Inches(0.34),
        title, 15, NAVY, bold=True)
    tb = s.shapes.add_textbox(x + Inches(0.24), Inches(5.12),
                              Inches(5.55), Inches(1.35))
    tf = tb.text_frame; tf.word_wrap = True
    for j, ln in enumerate(lines):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        r = p.add_run(); r.text = ln
        r.font.size = Pt(11.5); r.font.color.rgb = INK; r.font.name = 'Arial'

deploy_card(Inches(0.55), NAVY, 'NXP internal \u2014 3 stations', [
    'Romania \u00b7 Mexico \u00b7 China \u2014 one station each',
    'Used for platform & project development verification',
    'Shared YAML projects & benchmark keep all sites aligned',
])
deploy_card(Inches(6.78), ACCENT, 'Supplier factories \u2014 scale by demand', [
    'Suppliers decide station count per order volume',
    'Investment based on our published BoM list',
    'Acceptance runs against our benchmark station',
])

rect(s, Inches(0.55), Inches(6.82), Inches(12.23), Inches(0.4),
     RGBColor(0xE6, 0xF7, 0xF4), line=ACCENT, line_w=Pt(1))
txt(s, Inches(0.75), Inches(6.86), Inches(11.9), Inches(0.32),
    'Reusable: instruments & PC \u00b7 Custom per project: fixture & cabling (incl. assembly labor) \u00b7 PSU optional (DUT adapter first)',
    11, NAVY, bold=True)

# ================================================================ Slide 6
s = prs.slides.add_slide(BLANK)
header(s, 'MTK vs. Traditional Manual Testing',
       'Same board \u2014 very different test depth.')
RED_X  = RGBColor(0xB3, 0x3A, 0x3A)
GREEN_OK = RGBColor(0x00, 0x8A, 0x76)

caps = [
    ('R-GND & voltage',
     'Handheld DMM, one point at a time \u2014 slow and operator-dependent',
     '80 points scanned automatically in seconds, \u00b10.1% voltage accuracy'),
    ('Clock verification',
     'Oscilloscope eyeballing \u2014 \u201cclock exists\u201d, no precision guarantee',
     'Counters measure each clock to spec: 32.768 kHz / 4 MHz / 6 MHz'),
    ('Power-up sequence waveform',
     'Not possible \u2014 scopes offer at most 4 channels',
     '12 power rails captured at ms level; rise, overshoot & ripple auto-graded'),
    ('RF function tests',
     'Skipped \u2014 far too complex to operate by hand',
     'Wi-Fi, Bluetooth & GNSS tested automatically with pass/fail verdicts'),
]
cw2, ch2, cgap2 = Inches(6.0), Inches(2.0), Inches(0.23)
for i, (title, manual, mtk) in enumerate(caps):
    col, row = i % 2, i // 2
    x = Inches(0.55) + col * (cw2 + cgap2)
    y = Inches(1.6) + row * (ch2 + Inches(0.15))
    rect(s, x, y, cw2, ch2, BOX_BG, line=BOX_LN, line_w=Pt(1))
    rect(s, x, y, Inches(0.07), ch2, NAVY)
    txt(s, x + Inches(0.24), y + Inches(0.14), cw2 - Inches(0.48), Inches(0.36),
        title, 16, NAVY, bold=True)
    txt(s, x + Inches(0.24), y + Inches(0.6), cw2 - Inches(0.48), Inches(0.24),
        '\u2717  MANUAL', 11, RED_X, bold=True)
    txt(s, x + Inches(0.24), y + Inches(0.86), cw2 - Inches(0.48), Inches(0.5),
        manual, 12, MUTED)
    rect(s, x + Inches(0.24), y + Inches(1.38), cw2 - Inches(0.48), Pt(1), BOX_LN)
    txt(s, x + Inches(0.24), y + Inches(1.48), cw2 - Inches(0.48), Inches(0.24),
        '\u2713  MTK', 11, GREEN_OK, bold=True)
    txt(s, x + Inches(0.24), y + Inches(1.74), cw2 - Inches(0.48), Inches(0.26),
        mtk, 12, INK)

# 5th contrast card: full-width strip
y5 = Inches(5.9)
x5, w5, h5 = Inches(0.55), Inches(12.23), Inches(0.95)
rect(s, x5, y5, w5, h5, BOX_BG, line=BOX_LN, line_w=Pt(1))
rect(s, x5, y5, Inches(0.07), h5, NAVY)
txt(s, x5 + Inches(0.24), y5 + Inches(0.3), Inches(2.6), Inches(0.4),
    'Power net & test point selection', 15, NAVY, bold=True)
txt(s, x5 + Inches(3.1), y5 + Inches(0.12), Inches(4.4), Inches(0.24),
    '\u2717  MANUAL', 11, RED_X, bold=True)
txt(s, x5 + Inches(3.1), y5 + Inches(0.38), Inches(4.4), Inches(0.5),
    'Workers study design documents to locate power nets and test points',
    11.5, MUTED)
txt(s, x5 + Inches(7.85), y5 + Inches(0.12), Inches(4.1), Inches(0.24),
    '\u2713  MTK', 11, GREEN_OK, bold=True)
txt(s, x5 + Inches(7.85), y5 + Inches(0.38), Inches(4.1), Inches(0.5),
    'AI algorithm auto-selects power nets and the best test points',
    11.5, INK)

rect(s, Inches(0.55), Inches(6.98), Inches(12.23), Inches(0.4),
     RGBColor(0xE6, 0xF7, 0xF4), line=ACCENT, line_w=Pt(1))
txt(s, Inches(0.75), Inches(7.02), Inches(11.8), Inches(0.32),
    'Takeaway: MTK turns today\u2019s spot checks into full-coverage, high-precision, automated testing.',
    13, NAVY, bold=True)

# ================================================================ Slide 7
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
PANEL = RGBColor(0x17, 0x33, 0x52)
rect(s, Inches(0.55), Inches(0.42), Inches(0.12), Inches(0.62), ACCENT)
txt(s, Inches(0.85), Inches(0.36), Inches(11.9), Inches(0.75),
    'Summary & Ask', 30, WHITE, bold=True)
txt(s, Inches(0.87), Inches(1.02), Inches(11.9), Inches(0.4),
    'One decision unlocks real-mode validation and production rollout.', 14, LIGHT)
rect(s, Inches(0.55), Inches(1.42), Inches(12.23), Pt(1.4), RGBColor(0x2A, 0x45, 0x63))

sums = [
    ('All-in-one, ready to scale',
     'From design to platform, from net to system \u2014 one common platform for all NXP Commercial Boards.'),
    ('Professional automation, zero license fees',
     '\u2248 $15\u201319k one-time per station; no recurring software cost \u2014 built by NXP, for NXP boards.'),
    ('Proven, not paperware',
     'FCT accepted on a real board with no obvious bugs; ICT test plans built for 8 board types.'),
]
for i, (t, d) in enumerate(sums):
    y = Inches(1.75) + i * Inches(1.55)
    rect(s, Inches(0.55), y, Inches(6.85), Inches(1.4), PANEL)
    rect(s, Inches(0.55), y, Inches(0.07), Inches(1.4), ACCENT)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), y + Inches(0.22),
                              Inches(0.44), Inches(0.44))
    circ.fill.solid(); circ.fill.fore_color.rgb = ACCENT
    circ.line.fill.background(); circ.shadow.inherit = False
    txt(s, Inches(0.8), y + Inches(0.27), Inches(0.44), Inches(0.36),
        str(i + 1), 15, NAVY, bold=True, align=PP_ALIGN.CENTER)
    txt(s, Inches(1.42), y + Inches(0.16), Inches(5.85), Inches(0.34),
        t, 16, WHITE, bold=True)
    txt(s, Inches(1.42), y + Inches(0.56), Inches(5.85), Inches(0.75),
        d, 12, LIGHT)

rect(s, Inches(7.75), Inches(1.75), Inches(5.03), Inches(4.7),
     PANEL, line=ACCENT, line_w=Pt(1.5))
txt(s, Inches(8.05), Inches(1.95), Inches(4.4), Inches(0.4),
    'THE ASK', 18, ACCENT, bold=True)
asks = [
    ('Approve the instrument rack',
     '\u2248 $15.4\u201319.3k per station (N5747A optional) \u2014 unlocks Phase 3: real-mode, full-flow validation.'),
    ('Endorse the pilot rollout',
     'First production-floor station, onboarding more board types with the HWS-CC TE team.'),
]
for i, (t, d) in enumerate(asks):
    y = Inches(2.55) + i * Inches(1.75)
    txt(s, Inches(8.05), y, Inches(4.45), Inches(0.34),
        '\u25b6  ' + t, 14.5, WHITE, bold=True)
    txt(s, Inches(8.3), y + Inches(0.4), Inches(4.25), Inches(1.1),
        d, 12, LIGHT)

rect(s, Inches(0.55), Inches(6.85), Inches(12.23), Pt(1.2), ACCENT)
txt(s, Inches(0.87), Inches(6.95), Inches(11.9), Inches(0.4),
    'Developed & deployed by NXP HWS-CC Team \u2014 all TE engineers  \u00b7  Thank you',
    13, LIGHT)

os.makedirs('out', exist_ok=True)
prs.save('MTK_Common_Platform_Intro.pptx')
print('saved', len(prs.slides._sldIdLst), 'slides')
